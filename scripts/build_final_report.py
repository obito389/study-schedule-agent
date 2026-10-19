"""Build the course report from the retained teacher template and measured project evidence."""
import copy, csv, hashlib, json, pathlib, re, subprocess, xml.etree.ElementTree as ET
from collections import defaultdict
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
from report_figures import figures

ROOT=pathlib.Path(__file__).resolve().parents[1]
MAT=ROOT/'delivery/materials'; OUT=ROOT/'delivery/reports'; EV=ROOT/'delivery/evidence'; OUT.mkdir(parents=True,exist_ok=True)
template=next(MAT.glob('《*.docx')); doc=Document(template)
figures(ROOT)
roster=json.loads((MAT/'roster.json').read_text(encoding='utf-8-sig'))
commit=subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip()
repo='https://github.com/obito389/study-schedule-agent'
evidence={'coverage':{},'focused':{},'testSuites':[],'tables':{}}
for module in ['backend','backend-hjh']:
    xml=ET.parse(ROOT/module/'target/site/jacoco/jacoco.xml').getroot()
    evidence['coverage'][module]={n.attrib['type']:{'covered':int(n.attrib['covered']),'missed':int(n.attrib['missed'])} for n in xml.findall('counter')}
    for c in xml.findall('.//class'):
        if c.attrib['name'].split('/')[-1] in ['DeliveryCommands','DeliveryController','ActionExecutor','UndoExecutor','TeamSharedValidator','TeamActionService','GreedyScheduler','HardConstraintValidator','AgentService','ProfileService']:
            evidence['focused'][c.attrib['name']]={n.attrib['type']:{'covered':int(n.attrib['covered']),'missed':int(n.attrib['missed'])} for n in c.findall('counter')}
    for p in (ROOT/module/'target/surefire-reports').glob('TEST-*.xml'):
        root=ET.parse(p).getroot();evidence['testSuites'].append(dict(root.attrib))
schema=defaultdict(list)
for row in csv.reader((ROOT/'.runtime/schema.tsv').read_text(encoding='utf-8-sig').splitlines(),delimiter='\t'):
    if row and row[0]!='TABLE_NAME':schema[row[0]].append(row[1:])
evidence['tables']=dict(schema)
(EV/'whitebox-and-schema.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf8')
black=json.loads((EV/'blackbox.json').read_text(encoding='utf8'));performance=json.loads((EV/'performance.json').read_text(encoding='utf8'))
browser_text=(EV/'browser-tests.json').read_text(encoding='utf-8-sig');browser=json.loads(browser_text[browser_text.find('{'):])
if browser['stats']['unexpected']:raise RuntimeError('Browser acceptance has failures; report cannot claim completion')
test_count=sum(int(s['tests']) for s in evidence['testSuites'])

# Distill cover and page geometry before replacing the teacher's content slots.
normal=doc.styles['Normal'];normal.font.name='宋体';normal.font.size=Pt(12)
normal.element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
for name in ['Heading 1','Heading 2','Heading 3']:
    if name not in doc.styles:
        st=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH);st.base_style=normal
        level=OxmlElement('w:outlineLvl');level.set(qn('w:val'),str(int(name[-1])-1));st.element.get_or_add_pPr().append(level)
    st=doc.styles[name];st.font.name='黑体';st.font.size=Pt(14);st.font.bold=True;st.font.color.rgb=RGBColor(0,0,0)
    st.element.rPr.rFonts.set(qn('w:eastAsia'),'黑体')
for p in list(doc.paragraphs)[15:]:p._element.getparent().remove(p._element)
doc.paragraphs[6].text='学习日程 Agent 系统设计报告'
doc.paragraphs[6].alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.paragraphs[6].runs[0].font.size=Pt(16);doc.paragraphs[6].runs[0].bold=True
doc.paragraphs[11].text='年级：2024 级 2 班';doc.paragraphs[14].text='时间：2026 年 10 月 20 日'
for p in [doc.paragraphs[11],doc.paragraphs[14]]:p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.runs[0].font.size=Pt(15)
for name,number in roster.items():
    p=doc.add_paragraph(('组长：' if name=='熊翔' else '成员：')+name+'    学号：'+number);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.runs[0].font.size=Pt(12)
body_start=len(doc.paragraphs);counter=0;fig=0;units=[];used=set()
source_docs={name:Document(next(MAT.glob(pattern))) for name,pattern in [('黄健华','*后端开发工作报告*'),('赵欢','*个人项目报告*')]}
source_paragraphs={name:[p.text.strip() for p in source.paragraphs if len(p.text.strip())>90 and len(p.text)<1100 and len(re.findall('[\u4e00-\u9fff]',p.text))/max(1,len(p.text))>.40] for name,source in source_docs.items()}

def para(text):
    p=doc.add_paragraph(str(text));p.paragraph_format.line_spacing=1.25;p.paragraph_format.space_after=Pt(6);p.paragraph_format.first_line_indent=Pt(24)
    if str(text).startswith('测试源码位置') or 'https://' in str(text):
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.first_line_indent=Pt(0)
        for r in p.runs:r.font.size=Pt(10)
    return p
def table(title,headers,rows,note='本表按最终交付版实现和当前测试证据整理。'):
    global counter
    counter+=1;p=doc.add_paragraph(f'表 {counter}  {title}');p.paragraph_format.keep_with_next=True;p.runs[0].bold=True
    p=doc.add_paragraph('表注：'+note);p.paragraph_format.keep_with_next=True;p.runs[0].font.size=Pt(10)
    t=doc.add_table(rows=1,cols=len(headers))
    t.autofit=False
    widths=([1.35,4.35] if len(headers)==2 else [3.7,.7,1.3] if headers[0]=='套件' else [1.3,2.2,2.2] if len(headers)==3 else [1.4,1.1,.8,2.4] if headers[0]=='字段' else [1.5,1.0,1.0,2.2])
    for col,width in zip(t.columns,widths):col.width=Inches(width)
    if 'Table Grid' in doc.styles:t.style='Table Grid'
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');borders.append(b)
    t._tbl.tblPr.append(borders)
    for c,h in zip(t.rows[0].cells,headers):c.text=str(h)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    for row in t.rows:
        for cell,width in zip(row.cells,widths):
            cell.width=Inches(width)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.05
                for r in p.runs:r.font.size=Pt(10);r.font.name='宋体'
    return t
def image(file,title):
    global fig
    file=pathlib.Path(file)
    if not file.exists():return
    with Image.open(file) as im:w,h=im.size
    width=min(5.7,4.9*w/h);doc.add_picture(str(file),width=Inches(width));doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    fig+=1;p=doc.add_paragraph(f'图 {fig}  {title}');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.runs[0].font.size=Pt(10)
def page(title,texts=(),chapter=None):
    units.append(title)
    if chapter:
        h=doc.add_paragraph(chapter,'Heading 1');h.paragraph_format.page_break_before=True
    h=doc.add_paragraph(title,'Heading 2');h.paragraph_format.page_break_before=not bool(chapter)
    for text in texts:para(text)
def source(author,keyword):
    for i,text in enumerate(source_paragraphs[author]):
        if keyword in text and (author,i) not in used:
            used.add((author,i));para(f'阶段设计材料摘录（{author}）：'+text[:650]);para('以上摘录用于保留成员的设计思路和阶段贡献。涉及待联调、Mock、旧端口及旧测试数量的表述属于原报告当时状态，最终状态以本报告实现与测试章节为准。');return
def coverage_rows(data):
    rows=[]
    for k in ['INSTRUCTION','BRANCH','LINE','METHOD']:
        n=data[k];total=n['covered']+n['missed'];rows.append([k,n['covered'],n['missed'],f"{n['covered']/total*100:.2f}%" if total else '不适用'])
    return rows
def diagram(name,title,steps):
    path=EV/'screenshots'/name;im=Image.new('RGB',(1200,180+len(steps)*145),'white');d=ImageDraw.Draw(im);font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',31);small=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',25)
    d.text((50,25),title,font=font,fill='#163d3c')
    for i,(label,desc) in enumerate(steps):
        y=105+i*145;d.rounded_rectangle((55,y,1145,y+103),radius=12,fill='#eff6f3',outline='#71978c',width=2);d.text((80,y+12),label,font=font,fill='#143c3a');d.text((80,y+59),desc,font=small,fill='#46595d')
        if i<len(steps)-1:d.line((600,y+105,600,y+140),fill='#71978c',width=4)
    im.save(path);return path

page('交付摘要与版本说明',[
 '本报告记录学习日程 Agent 的最终集成成果。系统采用组员真实后端模块作为工程基础，接入组长经多轮调整的前端整体设计，实现任务管理、日历展示、规则排程、模型建议、人工校对导入、执行撤销、耗时反馈和偏好管理。最终版不是四份独立程序的展示合集，而是由同一数据库、同一日程版本和同一执行器约束的一条业务链路。',
 '报告依据仓库源代码、重新运行的测试和当前页面截图编写。两份成员个人报告作为署名过程材料保留，其历史测试记录不替代最终验收。图片和录音未接入 OCR 或 ASR；正式账户认证、生产部署和长期压力测试也不在本次已完成范围内。',
 f'代码基准提交为 {commit[:12]}，分支为 codex/final-delivery。后端默认 8087、前端默认 5181，使用独立 MySQL 数据库 final_delivery，测试服务和测试库另行隔离。'],chapter='交付说明')
table('交付材料与复现入口',['材料','内容','位置'],[['系统源码','Java、Vue、迁移、脚本和测试','本压缩包源码目录'],['团队报告','老师模板结构与成员贡献说明','delivery/reports'],['测试证据','黑盒、白盒、浏览器、性能','delivery/evidence'],['个人材料','两份原始署名报告及组长旧稿','delivery/materials']])
page('目录',chapter='目录')
p=doc.add_paragraph();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u');p._p.append(field)

page('课程设计的目的与要求',[
 '课程实践要求从一个明确的学生学习场景出发，完成需求分析、面向对象建模、数据库设计、详细设计、系统实现和可重复测试。项目以有截止时间、有课程占用、有个人偏好的任务安排为对象，训练团队把自然语言需求转成可校验的数据与动作，而不是把模型回答直接当作执行结果。',
 '学生往往同时面对课程、实验报告和复习任务，信息分散在聊天截图、语音记录和个人备忘中。系统需要让用户看见待办及时间约束，核对模型建议，并在确认后保存。课程实践的重点是可解释、可回退、可复现：输入信息不充分时要求补充，安排不满足约束时明确说明，重复请求不能产生重复任务。',
 '最终交付遵循老师追加的要求：封面标注组员姓名学号，报告超过七十页，所有表格之前给出编号和表注；测试包括黑盒、白盒及集成、并发、兼容性等验证；真实前端截图附说明；代码正文只保留少量关键表达式，完整源码提供链接和压缩包。'],chapter='一、课程设计的目的与要求')
table('课程要求与实现验收',['要求','本项目对应工作'],[['需求与建模','用例、类关系、时序和模块边界'],['数据库设计','16 张业务表与 Flyway 独立迁移'],['系统实现','前端到执行器再到 MySQL 的真实链路'],['测试与复现','自动化断言、当前截图和可重复命令']])

page('2.1.1 题目与问题界定',[
 '题目为学习日程 Agent 系统的分析、设计与实现。系统服务于需要协调课程、作业和复习任务的学生。这里的 Agent 负责读取上下文、理解要求和生成结构化建议；日程的最终写入由受控服务完成。用户确认是建议与正式数据之间的边界。',
 '任务记录与时间安排是不同对象。创建一个预计六十分钟、明晚截止的任务，并不代表它已占用日历中的某一小时；规则排程先生成候选时间块，确认后才落库。这个区分贯穿接口、页面状态、测试和数据库设计，避免界面显示成功但实际没有安排的歧义。',
 '本报告采用团队最终版作为统一基准。原完整版与组员联调版保留，用于回看界面设计和模块来源；报告中不混用三个版本的截图、接口与测试数字。'],chapter='二、设计报告正文\n2.1 概述')
page('2.1.2 主要目标与完成边界',[
 '主要目标包括任务新增修改与状态管理、课程及不可用时段、日周月日历、规则候选排程、真实 DeepSeek 建议、导入校对、执行记录及撤销、完成反馈和学习偏好。每一个可写操作都必须能说明输入是什么、哪些条件被校验、写入哪些对象以及如何处理失败。',
 '完成边界与产品目标分开表述。文本导入可以由模型提取字段，图片与录音只支持文件接收和人工校对。浏览器语音输入取决于浏览器能力；月历重要事项保存在当前浏览器，不宣称跨设备同步。身份选择沿用本地演示请求头，不作为正式用户认证。'])
table('最终功能状态',['功能','状态','验收依据'],[['规则排程与确认','已实现','候选不落库、确认与撤销测试'],['真实模型创建与查询','已接入','live-model.json 和当前截图'],['MySQL 与执行审计','已实现','集成测试与数据库实查'],['OCR / ASR','未接入','页面明确提示人工校对'],['正式登录与云同步','未实现','列入后续工作']])
page('2.1.3 开发与运行环境',[
 '后端以 Java 21 为编译目标，采用 Spring Boot 3.5.16、Spring MVC、Validation、Spring Data JPA 和 MySQL 8。Maven reactor 将 backend 核心模块作为依赖，backend-hjh 是启动服务。前端采用 Vue 3、TypeScript、Vite、Element Plus 和 FullCalendar，保持组员技术栈。',
 '本轮本机测试使用 Windows 和 Java 25 运行环境，编译产物仍为 Java 21 字节码。自动化验证注入固定 Clock，不把真实模型的不确定输出作为测试断言。前端生产构建和开发预览分开，依赖版本锁定在 package-lock.json 中。'])
table('开发与验证环境',['部分','配置'],[['数据库','MySQL 8.0.46，独立 final_delivery 系列库'],['后端','Java 21 target；Spring Boot 3.5.16'],['前端','Vue 3 / TypeScript / Vite / Element Plus'],['验证','JUnit 5、Mockito、MockMvc、JaCoCo、Vitest、Playwright'],['时间','Asia/Shanghai；接口 Instant 与存储本地时间显式转换']])

concepts=[
('2.2.1 用户分析','学生需要同时看见任务的截止、剩余耗时和已排时间。最常见的入口是手工新建、导入通知和自然语言要求；最常见的决策是确认建议、调整安排或取消任务。界面应把状态变化和失败原因展示出来，不能将待确认和已执行混为一谈。','用户','赵欢'),
('2.2.2 用例图与系统边界','学生是应用的主要参与者，模型提供方和数据库是外部依赖。模型仅参与识别和建议，不拥有修改日程的权限。创建任务、生成候选、确认写入、完成反馈及撤销形成不同用例，上传与人工校对也需要单独表示。','模块','赵欢'),
('2.2.2 各模块用例与异常场景','需求分析同时考虑正常路径与异常路径：耗时缺失需要澄清，版本过期需要重新预览，文件过大需要拒绝，模型中断不能改变日程。是否可撤销取决于执行类型及后续修改；完成反馈的证据不通过普通撤销入口删除。','状态','赵欢'),
('2.2.3 类图与领域关系','ScheduleSnapshot 是跨模块读取契约，TaskSnapshot 包含已排时间块；ActionSpec 描述候选动作，ImportProposal 描述待核对任务，ProfileFact 描述带证据的事实。数据库实体使用组员已有 Long 主键，适配层负责转成契约字符串 ID，不改变契约含义。','接口','黄健华'),
('2.2.4 时序图与确认边界','一次模型建议先读取快照，在数据库事务之外请求模型，再将合法候选保存为草案。确认时锁定当前用户并检查版本，重新校验约束后写入任务或时间块。返回成功后前端刷新快照；网络重试继续使用原操作键，防止重复执行。','事务','黄健华'),
('2.2.5 非功能需求与验收条件','需求中的可靠性表现为原子写入、重复请求不增写和失败后可恢复。可维护性表现为模块端口、固定测试输入和清楚的源码归属。性能验收限定为本机小样本和并发冒烟，不以几次请求推断生产容量，也不对未实现的认证能力作承诺。','测试','赵欢'),
('2.3.1.1 总体功能设计','系统采用模块化单体。Vue 页面调用最终版界面契约，适配层调用主线的 Agent、排程、画像与导入服务，再进入组员执行器、撤销器和 Repository。所有模块运行在一个 Spring Boot 进程中，避免为了课程实践引入独立微服务和消息系统。','分层','黄健华'),
('2.3.1.2 各模块功能设计','前端负责收集输入、展示候选和确认反馈；导入模块负责来源与状态机；Agent 负责受控工具选择；排程模块负责候选生成及硬约束；画像模块负责事实和耗时样本；执行器负责正式任务变更；持久化层负责关系约束和审计保存。','交接','黄健华'),
('2.3.2 数据流和稳定契约','四个稳定契约在模块之间传递数据。适配层处理接口响应格式、ID 类型、时区和版本差异；它不引入与组员实体平行的第二套正式任务数据库。Receipt 是界面的执行回执，保存于现有 plan_draft 的扩展 JSON 中，并关联实际动作审计。','ImportProposal','赵欢'),
('2.3.3 状态与错误设计','页面区分加载、待补充、待确认、已执行和失败。模型失败应显示未保存及重试入口；版本冲突应明确要求刷新。HTTP 状态与业务错误保持一致，未知服务器错误不直接暴露数据库异常。导入失败保留可重试状态，确认后的任务通过真实接口重新查询。','异常','黄健华'),
('2.3.4 版本和幂等总体策略','planning_revision 是同一用户日程的并发版本；用户行锁串行化写入，baseRevision 拒绝过时请求。operationKey 绑定请求内容，完全相同的重试返回原回执，换内容复用键则拒绝。草案确认读取服务端保存的候选，不采纳客户端临时替换的标题或时间。','幂等','黄健华'),
]
for i,(title,text,key,author) in enumerate(concepts):
    chapter='2.2 系统需求分析' if i==0 else '2.3 系统总体设计' if i==6 else None
    page(title,[text],chapter)
    if i in [1,3,4]:
        image(EV/'screenshots'/({1:'uml-usecase.png',3:'uml-classes.png',4:'uml-sequence.png'}[i]),title)
    elif i==6:
        steps=[('用户输入','手工任务、自然语言、上传文件'),('候选与校验','契约、约束、状态和版本检查'),('用户确认','尚未执行的建议由用户核对'),('事务落库','组员执行器、审计、MySQL 和新快照')]
        image(diagram(f'architecture-{i}.png',title,steps),'最终系统模块和确认边界')
    else:
        concept_rows={
          0:[['任务视角','查看截止、预计耗时、状态与所属课程'],['安排视角','查看课程占用、正式时间块和冲突'],['决策视角','核对候选后确认，失败时保留输入'],['反馈视角','记录实际耗时，纠正无效学习事实']],
          2:[['缺少耗时','返回待补充，不创建正式任务'],['版本过期','拒绝旧候选，读取新快照再生成'],['上游失败','显示失败提示，保持正式日程不变'],['后续引用','撤销前检查步骤和时间块关联']],
          5:[['一致性','并发两个旧版本请求仅一个成功'],['可恢复性','事务失败后任务数量和版本不变'],['可测试性','固定时钟、隔离数据库和模型夹具'],['性能范围','本机短样本与并发冒烟，不推断生产容量']],
          7:[['Agent 与导入','输出结构化候选，保留来源及待补充项'],['排程与校验','计算可行时间块，独立复核硬约束'],['执行与撤销','锁用户、检查版本、原子写入及逆操作'],['画像与统计','记录有效样本，排除已否认或删除事实']],
          8:[['ScheduleSnapshot','读取时区、当前时刻、任务、课程和版本'],['ActionSpec','描述动作、参数、权限范围与约束提示'],['ImportProposal','保存来源、提取字段、问题及确认状态'],['ProfileFact','表达键值、证据、状态与事实版本']],
          9:[['待补充','问题与输入并列展示，允许修改再提交'],['待确认','展示候选，不把预览标为执行成功'],['失败与冲突','可理解提示，保留重试或刷新入口'],['已执行','展示回执并刷新真实数据快照']],
          10:[['版本基线','所有写入携带当前 planningRevision'],['锁与重检','用户行锁内再次校验归属、版本和约束'],['同键同内容','复用已保存回执，不重复产生任务'],['同键换内容','拒绝冲突，要求调用方使用新操作键']]
        }
        table(title+'的设计要点',['对象','责任与校验'],concept_rows[i])
        if i==7:
            para('模块交付以跨模块可调用为标准。前端通过 API 客户端统一身份头和响应解包；适配层将界面命令转换为已有服务调用，正式任务与时间块仍由组员执行器保存。导入确认、规则候选确认和 AI 建议确认虽然入口不同，最终都必须遵守相同的用户归属、版本和硬约束。')
            para('组长负责模块边界、数据契约及端到端验收，不能用单个模块的演示成功替代整条链路通过。成员原始代码和署名报告保留，集成修复在独立分支记录，便于分别说明原始贡献和最终联调工作。')
    source(author,key)

page('2.4.1 数据库表与建模原则',[
 '最终版沿用黄健华设计的十六张业务表。账户是归属根，学期组织课程，课程节次表达重复上课时间，任务和步骤表达目标分解，时间块表达正式安排。导入、草案、动作、执行日志、事实、事件和统计分别保存来源、候选、审计和反馈，避免把所有业务字段塞入一个 JSON。',
 '数据库结构由 Flyway 管理。组员原始迁移保留；最终版使用独立迁移副本，仅去掉写死的建库与切库语句，保证迁移只作用于部署时指定的数据库。契约扩展 JSON 用于完整保留 DTO，核心字段仍有明确列和关系约束。'],chapter='2.4 数据库设计')
table('业务表分组',['分组','表名'],[['账户与课程','app_user、semester、course、course_session'],['偏好与时间','user_preference、availability_window'],['任务执行','task、task_step、time_block'],['建议与审计','import_job、plan_draft、agent_action'],['反馈学习','execution_log、profile_fact、profile_event、learning_stat']])
page('2.4.2 数据表之间的关系',[
 'task.user_id 和 time_block.user_id 直接表达用户归属。course 通过 semester 归属于用户，task_step 通过 task 归属；因此接口不能只检查某个 ID 是否存在，还必须沿关系确认它属于当前用户。时间块中的 step_id 需要属于同一个任务，来源草案与导入任务也需要校验归属。',
 '撤销创建父对象时，后续新增的子步骤或时间块可能仍引用它。最终版在删除式逆操作之前检查这些后续引用，避免数据库外键错误或悄悄删除后来产生的数据。执行与撤销都必须重新检查硬约束，版本号只能保护并发顺序，不能替代业务约束。'])
image(EV/'screenshots/er-overview.png','当前 16 表外键关系概览，箭头由被引用表指向引用表')
table_info={
 'app_user':('用户根与日程版本','保存身份选择、时区、日程版本和画像版本。每次正式写入先锁定此行；用户归属与乐观版本检查由这条记录提供共同入口。'),
 'semester':('学期','学期开始日期与教学周数是课程节次展开的基础。最终界面单次课程用一周周期表达，原有单双周数据仍可由组员实体读取。'),
 'course':('课程','课程关联学期，提供任务所属课程的名称。删除前检查关联任务，避免外键被破坏；当前界面新增课程通过版本服务完成。'),
 'course_session':('课程节次','节次记录星期、开始结束时间、起止周和单双周规则。快照把重复规则展开成可比较的课程时间块，排程与执行使用同一占用集合。'),
 'user_preference':('用户偏好参数','组员原有偏好表作为领域基础保留。最终工作台的每日上限通过带证据的 ProfileFact 服务管理，不允许旧 CRUD 绕过版本校验直接修改。'),
 'availability_window':('可用或禁止时段','记录按日期或星期表达的窗口。最终界面新增不可用时段会与已排任务进行冲突检查，失败时事务整体回滚。'),
 'task':('任务','任务保存标题、截止时间、优先级、剩余耗时和状态。创建任务不隐含排程，完成时将剩余耗时设为零并记录反馈，取消时释放已排时段。'),
 'task_step':('任务步骤','步骤通过 task_id 归属于用户。组员执行器支持步骤操作，当前前端未提供单独的步骤编辑器，不能把接口能力等同于完整界面功能。'),
 'time_block':('正式时间块','保存已确认的任务占用、锁定状态和来源草案。只有 PLANNED 参与活动占用，完成与取消的时段不继续阻挡后续规则排程。'),
 'import_job':('导入任务','保存来源、状态、文件名、候选和契约数据。文字提取调用模型，图片及录音返回待人工填写的候选；确认提交通过执行器创建关联任务。'),
 'plan_draft':('服务端草案与回执','保存版本基线和候选数据。界面回执也保存在这张现有表中，以区分用户尚未确认的建议和已经完成的操作，不新增平行持久化体系。'),
 'agent_action':('动作审计','记录动作类型、前后快照、请求指纹、版本与所属草案。审计提供幂等依据和撤销分组；不能通过普通 CRUD 修改审计来伪造执行结果。'),
 'execution_log':('任务完成反馈','保存实际耗时和反馈键。重复反馈不会重复产生样本；已经完成的任务反馈不能在普通执行记录中随意撤销，以免学习证据与状态脱节。'),
 'profile_fact':('带证据的用户事实','事实包含键、值、来源、强度、状态、证据和版本。用户否认或删除会追加状态变化，历史仍保留，失效事实不参与当前估计。'),
 'profile_event':('画像事件','记录事实变化的事件与内容，形成可追溯的状态链。完成反馈、手动偏好及否认删除均能够找到相应事实和事件记录。'),
 'learning_stat':('耗时估计统计','统计样本数量和倍率以供解释。估计服务在样本不足时保持原值，否认样本后同步更新统计，避免页面显示已经失效的学习结果。')}
for table_index,(name,(label,description)) in enumerate(table_info.items()):
    rows=schema[name];page((f'2.4.3 数据库表结构 {name}' if table_index==0 else f'2.4.3.{table_index+1} {name} 表结构与约束'),[description,f'下表来自当前 MySQL 的 information_schema，而非仅根据设计稿推测。{name} 当前包含 {len(rows)} 个字段。字段是否可空、主键和索引标记按实际结构列出，具体关系与初始化过程可追溯到 db/team-migration。'])
    table(name+' 字段字典',['字段','类型','空值/键','业务说明'],[[r[0],r[1],r[2]+('/'+r[3] if r[3] else ''),r[4] or ('主键或关系标识' if r[0].endswith('id') else '以实体定义与服务校验为准')] for r in rows],note='来源为 final_delivery 当前库结构；YES 表示数据库允许 NULL，PRI/MUL/UNI 为 MySQL 键标记。')
    para('实现与验证：对应 business 包实体和 Repository，迁移在独立空库执行。通过快照和真实 MySQL 集成断言核对状态；归属、版本及业务范围由应用服务继续校验。')

screens=[('任务清单','01-tasks.png','任务列表将课程、截止、耗时和状态集中展示，支持搜索、筛选及排序。任务状态和已排时间块分开判断，已完成任务可以在相应筛选下查看。'),('新建与修改任务','02-task-editor.png','任务表单收集名称、北京时间截止、预计耗时、优先级和所属课程。预计耗时范围为 1 至 10080 分钟，优先级与组员后端统一为 1 至 3，标题限制为 100 字。'),('规则排程预览','03-plan-preview.png','候选列表说明任务、时间和缺口。生成预览不改变正式时间块；用户确认后才调用执行器。预览版本与当前日程不一致时禁止继续保存，需要重新生成。'),('工作日日历','04-calendar-week.png','日历保留组长多轮修改后的布局，小时刻度清楚，默认展示工作日并可单独切换周末。事件左侧显示时间、右侧显示事项；日期分隔与表头对齐，取消内部滚动条造成的局部遮挡。'),('月历与重要事项','05-calendar-month.png','月视图以有安排的日期作为重点，不把每日详细任务挤进狭小单元格。点击日期可填写重要事项；这类备注保存在当前浏览器，刷新可恢复，不表示服务端或跨设备同步。'),('AI 学习助手','06-ai.png','输入学习要求后显示处理中、待补充、待确认、已保存或失败。模型摘要不能直接视为已执行；确认按钮仅在有效候选且版本未过期时可用。页面明确说明本次要求和日程上下文会发送至模型。'),('导入与人工校对','07-import-review.png','导入页保留上传、状态、字段校对和确认入口。图片与录音不宣称已自动识别，用户填写或修正名称、截止和耗时后才能创建正式任务。失败时展示可恢复状态，避免将失败当成保存成功。'),('偏好与耗时反馈','08-profile.png','每日上限通过事实服务保存，修改不得使已有安排违反约束。耗时估计显示样本数、倍率和依据；用户可以否认或删除事实。完成任务收集实际分钟数，为后续估计提供真实反馈。'),('执行记录和撤销','09-history.png','执行记录按新旧顺序展示，回执关联真实草案和审计。撤销前提示影响，后续修改或约束变化可能导致拒绝；完成反馈不通过普通撤销入口回退。')]
for i,(title,file,explanation) in enumerate(screens):
    page(f'2.5.{i+1} {title}',[explanation],chapter='2.5 交互界面设计' if i==0 else None)
    image(EV/'screenshots'/file,'最终版实际页面 '+title)
    para('操作说明：'+explanation)
    para('截图环境为本机最终版前端 5181 与后端 8087。业务数据为演示输入；真实模型调用的独立记录保存在 live-model.json。页面显示成功后仍以刷新快照及数据库验证作为最终判据，图片本身不替代自动化断言。')

flows=[
('任务创建与修改','表单输入先通过页面约束，再由界面契约适配成组员 ActionCommand。请求携带操作键、版本和写入范围。标题、耗时、优先级、课程归属及截止时间必须满足服务校验，事务提交后刷新任务列表。','CREATE_TASK → Task → agent_action → Receipt','任务'),
('规则排程与缺口','排程以截止时间、优先级和 ID 稳定排序，按三十分钟时间槽搜索。每个候选使用共享硬约束检查。若某任务不能完整排入，不把不完整片段悄悄当成完成安排，而是报告缺口。确认过程重新检查版本。','ScheduleSnapshot → GreedyScheduler → candidate → confirm','排程'),
('模型建议与工具边界','模型调用在数据库事务外执行，返回受限 JSON。Agent 先处理歧义，再调用注册工具。模型指定的连续窗口和任务目标必须保留，最终适配层重新限制候选范围；指定的窗口不能被默认规则静默忽略。','instruction → ModelAdapter → ActionSpec → server draft','模型'),
('导入状态机与重试','导入请求键关联来源与候选，状态从处理中进入待校对或失败。用户确认的候选需要消除歧义并提供有效字段。提交失败后保留已确认内容用于原样重试；同一个导入确认不重复创建任务。','PROCESSING → REVIEW → COMMITTED / FAILED','确认'),
('画像事实与样本','用户显式偏好和任务完成样本来源不同。事实保存来源、证据、强度和版本；否认与删除追加状态变化。耗时估计样本少于三个时保持原值，足够样本采用平滑倍率，人工剩余耗时不被估计静默覆盖。','ExecutionLog → ProfileFact → ProfileEvent → LearningStat','耗时'),
('撤销与逆序恢复','撤销使用原动作前后快照，并按动作 ID 逆序恢复同一草案。撤销前检查当前状态是否仍与原执行结果匹配，同时检查关联对象和当前硬约束。创建任务并带时间块的动作必须先删除子块，再回退父任务。','lock → revision → guard → reverse actions → new revision','撤销'),
('课程与固定时段','课程通过学期和节次保留组员建模，固定时段使用 availability_window。界面传入北京时间后显式转换，新增时检查当前任务是否冲突。删除课程先检查是否仍有关联任务，失败时整体保持原状态。','semester / course_session / availability_window → snapshot','课程'),
('接口设计与响应适配','原组员接口采用 code/message/data，组长界面原先使用直接 DTO。最终版在客户端统一解包成功响应，保留 HTTP 失败状态；界面专用适配接口返回快照或回执，原组员执行器接口仍负责正式写入。','Vue → DeliveryController → Team services → member executor','接口'),
('时区和日历显示','接口传递 Instant，业务数据库使用显式转换后的 LocalDateTime。前端表单以北京时间输入，再转换为 UTC；日历将北京时间的浮动日期用于显示，避免浏览器处于不同时区时出现八小时偏差。测试覆盖跨午夜和无效时间。','Shanghai input ↔ Instant ↔ member LocalDateTime','数据')]
for i,(title,text,path,key) in enumerate(flows):
    label='2.6.1 业务流程设计 ' if i==0 else '2.6.2 接口设计 ' if i==7 else '2.6.3 ' if i==8 else f'2.6.1.{i+1} '
    page(label+title,[text,'处理路径：'+path],chapter='2.6 详细设计' if i==0 else None)
    table(title+'控制点',['阶段','正常结果','失败结果'],[['输入','字段和对象明确','缺失、格式或范围错误即拒绝'],['预览/校验','形成候选或可执行动作','保留原数据并展示原因'],['确认/事务','正式写入与版本推进','过期、冲突或越权不产生部分写入'],['重试/验证','同键同内容复用原结果','同键换内容拒绝']])
    source('赵欢' if i in [2,3] else '黄健华',key)

implementation=[('工程组织和启动入口','backend 提供稳定契约和核心能力；backend-hjh 提供实体、服务、执行器和启动类。Maven reactor 保证依赖按顺序编译，核心 jar 与可执行启动 jar 分开。前端使用 Vite 代理访问本地服务，不在页面中保存模型密钥。'),('事务、审计和持久化落地','最终版不把前端替换理解成后端重写。任务正式写入仍经组员 ActionExecutor，撤销仍经 UndoExecutor，画像和导入也落到原实体。DeliveryCommands 将界面回执与同一草案中的动作关联，便于历史查询、幂等与恢复。'),('前端设计迁移与适配','保留组长已有布局、设置入口、任务表单、AI 建议和日历细节，迁移时逐一检查每个 API 调用。优先级和字段长度与组员数据库约束一致；界面文案将图片录音明确标为人工校对，并修复模型失败提示和执行记录排序。'),('组长个人负责模块与实际贡献','熊翔原定负责架构、稳定契约、Agent、排程、用户模型和集成评审。最终阶段负责需求取舍、模块接口适配、回归与交付整理，并参与界面方案与交互优化。前端工程和组件原型来自樊华钰，组员后端和导入工作分别保留署名；本报告不将所有代码归为组长独立编写。实现和文档整理使用 AI 辅助，组长承担需求判断、修改与验证。'),('集成中发现的问题与修复','联调发现的缺口包括写死数据库名、响应包裹不同、原生 ID 与契约 ID 类型不同、完成时间块仍阻挡新排程、加载前误用旧版本、模型错误码直接显示和历史记录未明确排序。每项修复都围绕真实业务链路，并通过当前测试复核，不将旧报告中的计划当作已完成工作。')]
for i,(title,text) in enumerate(implementation):
    page(f'2.7.{i+1} {title}',[text],chapter='2.7 系统实现' if i==0 else None)
    if i==3:table('组内分工与实际贡献',['成员','原定职责','最终集成贡献'],[['熊翔','架构、Agent、排程、用户模型','集成、约束、测试、交付与界面优化'],['黄健华','数据库与事务后端','16 表、执行审计、幂等、撤销'],['赵欢','导入与模型适配','导入契约、状态机、校对和模块测试'],['樊华钰','前端开发','工程骨架、日历、状态及事实组件原型']])
    else:source('黄健华','实现' if i==0 else '测试' if i==4 else '事务')
    para('完整代码以固定提交链接为准：'+repo+'/tree/'+commit)

page('2.8.1 测试策略与环境隔离',[
 '测试分为内部白盒、外部 HTTP 黑盒、真实 MySQL 集成、浏览器端到端、并发与小样本性能、数据库持久化等层次。白盒依据代码路径和 JaCoCo 指标，黑盒只依据输入、HTTP 输出和可见业务状态。两者覆盖角度不同，不把通过的测试数量直接称为分支覆盖率。',
 '固定测试时间和本地模型提供方保证断言可重复；真实 DeepSeek 另外进行连通和确认落库检查。预览库、集成测试库与浏览器测试库彼此隔离，不对组长原完整版数据库执行清理或迁移。'],chapter='2.8 系统测试')
table('本轮验证汇总',['层次','执行结果','证据'],[['后端回归与集成',f'{test_count} 项；以 XML 报告为准','surefire 和 whitebox-and-schema.json'],['前端单元','17 项通过','Vitest 时间、筛选、日历工具测试'],['外部 HTTP 黑盒',f"{black['testsRun']} 项通过",'blackbox.json'],['浏览器主流程',f"{browser['stats']['expected']} 项通过",'browser-tests.json'],['小样本性能','130 个快照请求均返回 200','performance.json']])
case_descriptions=[
('健康快照和北京时间','GET /schedule','200；timezone=Asia/Shanghai','确认客户端读取的是同一后端快照及明确时区，而不是静态样例。'),
('缺失身份','不发送 X-User-Id','401，不能读取快照','身份缺失必须被拦截；本项验证归属入口，不等同于正式登录认证。'),
('创建请求幂等','同键、同版本、同 payload 连续创建两次','两次 200，正式任务仅一条','网络重试或重复点击不能重复创建任务；第二次返回原回执。'),
('同键换内容','创建成功后复用操作键修改标题','409，不能覆盖原任务','幂等键必须绑定内容，不能把不同请求误认为同一次重试。'),
('日程版本过期','版本推进后提交 baseRevision=0','409，拒绝过期修改','多人操作或多个页面造成过时输入时，应刷新并重新确认。'),
('耗时与优先级边界','durationMinutes=0；priority=5','400，任务数量不增加','分别检查整数范围的下界和上界；最终组员优先级范围为 1 至 3。'),
('过去截止和回滚','deadline=2020-01-01T00:00:00Z','4xx；版本仍为 0','拒绝无效截止时间，确认拒绝后未发生部分写入或版本推进。'),
('预览确认与撤销','创建任务→预览→确认→撤销','预览无块；确认有块；撤销清除块','验证整条跨模块链路，而不是只检查一个接口返回成功字符串。'),
('完成与画像反馈','完成任务并记录 actualMinutes=90','DONE、存在反馈事实、不可普通撤销','确认完成状态、实际耗时和学习证据相互一致。'),
('跨用户对象访问','其他用户引用原用户 taskId','403 或 404，原任务不变','对象归属不仅检查用户头，还必须检查目标任务所属用户。'),
('旧 CRUD 写入保护','POST /tasks 空请求','409，提示使用统一执行入口','避免旧控制器绕过版本、硬约束与审计直接改正式状态。'),
('偏好与版本联动','设置每日上限后重用旧版本','首次 200；旧版本 409','偏好也会影响排程，修改后旧建议必须失效。'),
('模型澄清和失败','固定模型输入缺少信息与服务失败','NEEDS_CLARIFICATION / FAILED；无新增任务','失败不能显示为已执行，用户应获得补充输入或重试指引。'),
('同版本并发竞争','两个线程用不同键和同一版本创建','一个 200、一个 409；任务一条','检验用户行锁与版本检查组合，防止两个并发写入都使用旧版本。'),
('课程实体写入','添加未来一小时课程','200，快照存在课程和节次','确认最终界面确实接到组员课程实体，而不是只在浏览器保存。')]
for i,(case,desc) in enumerate(zip(black['cases'],case_descriptions),1):
    title,input_text,expected,meaning=desc
    page(f'2.8.2 黑盒 BB{i:02d} {title}',[meaning,'本用例通过独立 HTTP 客户端访问专用测试服务，不读取 Java 私有方法作为判定依据。每个用例先创建独立测试用户，避免已有任务或偏好干扰，断言由 scripts/verify_delivery.py 中同名测试方法实现。'])
    table('黑盒用例 '+f'BB{i:02d}',['项目','记录'],[['测试方法',case['case']],['前置条件','独立用户、专用 final_delivery_e2e 数据库和 8088 测试服务'],['输入',input_text],['预期',expected],['实际',case['status']+'；运行时全部断言满足'],['执行耗时',str(case['milliseconds'])+' ms'],['判定','通过' if case['status']=='PASS' else '失败']])
    para('执行步骤：建立前置状态，发送上述输入，读取返回状态与业务快照，再比较预期。涉及写入的用例同时检查任务数量、状态、时间块或版本，避免只有 HTTP 200 但没有真实效果。原始执行记录与脚本一并提交，可在隔离测试库中重新运行。')
    para('结果解释：'+meaning+' 当前通过范围仅为表中给出的输入和断言，不扩展为对所有可能输入、无限并发或生产环境的保证。')

page('2.8.3 白盒覆盖率实测',[
 '使用 JaCoCo 0.8.14 在 Maven 测试阶段插桩，分别生成核心模块与后端模块的 XML 和 HTML 报告。指令、行、方法和分支是不同统计口径。下列分母包含该模块被报告的全部类，旧 CRUD 和未走到的保护路径也在内，未通过排除代码来提高数字。',
 '核心模块覆盖率较高，组员服务模块仍包含大量保留的 CRUD、实体访问器和异常分支，覆盖率较低。白盒结果因此作为当前已验证范围与剩余风险的量化说明，不声称达到百分之百覆盖。后续应优先补未覆盖的撤销保护、复杂批量变更及正式认证路径。'])
for module,data in evidence['coverage'].items():table(module+' 覆盖率',['口径','已覆盖','未覆盖','比例'],coverage_rows(data))
white=[('排程算法与稳定排序','GreedyScheduler','截止相同按优先级和 ID 排序；不足容量整任务缺口；重复输入结果确定。','backend/src/test/java/com/studyplanner/planning/GreedySchedulerTest.java'),('共享硬约束分支','HardConstraintValidator','截止、重叠、锁定、禁止时间及每日上限；同时覆盖跨午夜和边界相等情形。','backend/src/test/java/com/studyplanner/planning/HardConstraintValidatorTest.java'),('执行器和幂等分支','ActionExecutor','检查新请求、完整重试、同键换内容、过期版本、权限和批次内部重叠。预览不得修改实体。','backend-hjh/src/test/java/com/studyplanner/review/ExecutorReviewRegressionTest.java'),('撤销保护与界面回执','DeliveryCommands','完整回执重试、篡改操作键、执行后撤销、后续子对象阻挡和已完成反馈不可普通撤销。','backend-hjh/src/test/java/com/studyplanner/team/TeamAssemblyIntegrationTest.java')]
for i,(title,cls,text,file) in enumerate(white,1):
    page(f'2.8.3.{i} 白盒路径 {title}',[text,'路径选择先阅读对应实现的判断条件，再使用固定输入构造正常分支和拒绝分支。测试检查返回状态，也检查数据库或快照在拒绝后保持原样。该方法可以定位到具体控制流，补足仅从接口成功返回无法发现的内部状态问题。'])
    selected=next((v for k,v in evidence['focused'].items() if k.endswith('/'+cls)),None)
    if selected:table(cls+' 实测指标',['口径','已覆盖','未覆盖','比例'],coverage_rows(selected))
    table('白盒路径与判定',['路径','预期状态'],[['正常输入','进入候选或执行路径，结果可查询'],['边界输入','按定义接受或拒绝，不发生截断隐式修正'],['版本或权限拒绝','写入前停止，原状态不变'],['重复调用','不重复产生正式对象或反馈']])
    para('测试源码位置：'+file)

page('2.8.4 单元与真实数据库集成',[
 '核心回归包括契约序列化、模型失败和澄清、规则排程、硬约束、画像和耗时估计。组员执行器回归用于验证预览、幂等、权限及撤销镜像。新增集成测试在真实 MySQL 中运行，证明适配层、组员表、草案和审计确实发生关联。',
 '下表按当前 Surefire XML 汇总。跳过、失败和错误单独保留；不能把未配置数据库时跳过的集成测试计为通过。本轮显式提供专用测试库地址，所有集成用例均执行。'])
table('后端测试套件',['套件','测试数','失败/错误/跳过'],[[s['name'].split('.')[-1],s['tests'],'/'.join(s.get(k,'0') for k in ['failures','errors','skipped'])] for s in evidence['testSuites']])
page('2.8.5 浏览器端到端和窄屏兼容',[
 'Playwright 从最终前端入口操作真实组件，覆盖新建任务、规则候选确认、执行记录撤销、完成反馈、导入校对、偏好保存与刷新恢复，以及 AI 待确认、确认和上游失败提示。浏览器场景使用固定本地模型提供方，避免真实模型不确定输出造成断言漂移。',
 '窄屏场景使用 444×706 视口检查导航和月历，并断言文档宽度不超过窗口。另以 1440×1000 桌面视口检查主要布局。当前只记录已运行的 Chromium 自动化环境，未运行的 Safari、Firefox 和真机语音能力不列为通过。'])
table('浏览器测试',['场景','结果'],[['任务→确认→撤销→完成','通过'],['导入→偏好→刷新→窄屏月历','通过'],['AI 待确认→确认→服务失败','通过'],['页面脚本异常','捕获并要求空错误列表']])
page('2.8.5.1 窄屏布局实测说明',[
 '以下截图采用 444×706 的浏览器视口，从最终版 AI 页面入口重新加载后采集。导航在顶部横向排列，设置和执行记录置于工具栏，输入区及建议区按单列展示。截图对应交付版本的真实界面，不是界面设计稿。',
 '窄屏验收重点检查页面能否自然纵向展开、输入框及主要操作是否可见，以及日历月视图是否出现横向溢出。弹窗关闭后应恢复主体操作，执行记录抽屉不应遮挡重新进入的页面。当前结果只代表已运行的 Chromium 视口测试，不能替代不同手机系统上的真机验收。'])
image(EV/'screenshots/10-mobile.png','最终版窄屏布局')
page('2.8.6 并发、事务和失败恢复',[
 '并发测试让两个线程使用相同 baseRevision 和不同操作键创建任务。正确结果是一条写入成功，另一条由于版本过期被拒绝，正式任务数量保持一条。该测试同时验证数据库行锁和应用层版本协作，不能仅靠前端禁用按钮解决并发。',
 '事务回滚测试包括过去截止、课程或固定时段与已有安排冲突、每日上限使已有日程失效、跨用户访问及撤销冲突。失败后检查版本和正式对象数量，确保没有留下半个任务或半批时间块。模型失败发生在写入前，页面展示失败而不是成功解读。'])
table('失败恢复检查',['触发','恢复行为'],[['旧版本确认','刷新并重新生成建议'],['网络重试','保留操作键，同请求复用原结果'],['同键换内容','拒绝，要求新的操作键'],['导入提交失败','保留已确认候选，支持原样重试'],['撤销存在后续引用','拒绝，先处理后续操作']])
page('2.8.7 本机小样本性能',[
 '性能部分使用外部客户端读取日程快照，先顺序采样五十次，再以八个线程并行读取八十次。统计中位数和九十五分位延迟，并检查所有 HTTP 状态。测试运行在本机、同机 MySQL 和测试数据库上，结果仅用于发现明显的接口退化。',
 '本测试没有模拟大量历史任务、长期持续负载、跨网络延迟或多用户生产部署，不能将短时请求数换算为系统生产吞吐能力。报告保留原始统计和执行脚本，后续可在明确数据规模及硬件条件后扩展压力测试。'])
table('本轮性能采样',['指标','实际值'],[[k,v] for k,v in performance.items() if k!='note'])
page('2.8.8 数据持久化与真实模型连通',[
 '真实模型检查独立于固定模型测试。最终服务调用 DeepSeek 生成明晚复习数据库的创建建议，返回待确认状态后再确认；随后读取快照，检查任务和时间块。该操作验证真实凭据与接口连通，不把某次模型响应视为全部语义场景准确率。',
 '持久化检查核对接口快照与 MySQL 中的正式任务、已排时间块和日程版本。整机或服务重启后的恢复检查尚未完成，本轮服务重启操作被自动审批阻止，因此不计入已通过项。月历重要事项属于浏览器存储，应单独通过刷新与页面恢复验证，不能把它写成 MySQL 持久化。'])
live=json.loads((EV/'live-model.json').read_text(encoding='utf8'))
table('真实模型记录',['字段','实测值'],[['提供方',live['provider']],['返回状态',live['status']],['建议类型',live['action']['type']],['模型耗时',str(live['elapsedMillis'])+' ms'],['数据检查','确认后由接口快照及数据库实查核对']])

page('总结与后续工作',[
 '本项目将原本独立提交的核心、后端、导入与前端工程整理成统一可运行版本，并采用组长已确认的界面设计改善交互。完成的价值在于把任务、候选、确认、执行、审计、撤销和反馈连接起来，形成可查询、可测试、可解释的数据链路。',
 '组长负责架构、稳定契约、Agent、排程、用户模型与集成评审，额外参与界面需求和交互优化。黄健华、赵欢和樊华钰的模块来源在源码和本报告中保留，不以统一成稿掩盖分工，也不以旧报告的待办替代当前实际实现。',
 '仍需完善的工作包括正式认证、OCR/ASR、跨设备重要事项同步、步骤编辑界面、更多浏览器与真机验证，以及覆盖率不足的保护分支。后续开发应继续从明确的失败样例和真实约束出发，保持凭据在环境变量中、模型建议与正式写入分离。'],chapter='三、总结')
page('参考文献与项目材料',[
 '[1] 学习日程 Agent 项目源码，codex/final-delivery，固定提交见交付说明。',
 '[2] 黄健华．学习日程 Agent 后端开发工作报告．2026 年 10 月．原始署名文件随交付包保留。',
 '[3] 赵欢．学习日程 Agent 个人项目报告 25 页版．2026 年 10 月 18 日．原始署名文件随交付包保留。',
 '[4] 熊翔．学习日程 Agent 系统设计报告．2026 年 10 月 10 日．本报告更新其阶段状态与集成结果。',
 '[5] Java SE 21 API Specification，java.time.Clock；Spring Boot 3.5 Reference Documentation。用于核对项目所用时间抽象和工程组织。',
 '[6] JUnit 5 User Guide；JaCoCo Java Code Coverage Library；Playwright Documentation；Vue 3 Documentation。测试与界面工具以当前仓库锁定依赖为准。',
 '本报告的当前实现判断和测试数字来自随包提交的源码、测试 XML、覆盖率 XML、HTTP 黑盒记录及浏览器 JSON。历史成员材料用于设计过程与贡献追溯，不作为尚未运行能力的完成证明。'],chapter='四、参考文献')
page('源码链接与少量关键逻辑',[
 '完整代码通过固定提交链接及源码压缩包提供，正文不附整份类文件。以下只保留与设计决策直接相关的简短逻辑，便于答辩解释；实际边界和异常处理应阅读对应源码及测试。',
 '排程排序：截止时间升序 → 优先级降序 → 任务 ID。',
 '执行入口：锁定用户 → 幂等检查 → 检查 baseRevision → 硬约束 → 批量写入 → 推进版本。',
 '撤销入口：读取原审计 → 当前状态镜像校验 → 逆序执行反向动作 → 新版本与审计。',
 '画像估计：样本不足三个保持原耗时；充足样本做先验平滑，倍率限制在 0.5 至 2.0。',
 repo+'/tree/'+commit],chapter='五、附录')
table('关键源码入口',['位置','说明'],[['backend/.../planning','GreedyScheduler 与 HardConstraintValidator'],['backend/.../agent','AgentService 与受控工具'],['backend-hjh/.../executor','组员统一执行与请求指纹'],['backend-hjh/.../undo','组员快照逆序撤销'],['backend-hjh/.../delivery','最终界面契约与回执适配'],['frontend/src/pages','最终页面与日历'],['scripts/verify_delivery.py','HTTP 黑盒与性能复现']])
page('交付复现与五分钟讲解安排',[
 '先创建独立数据库并配置 TEAM_DB_URL、TEAM_DB_USER、TEAM_DB_PASSWORD，再运行 Maven 构建和启动服务。前端通过 npm ci 安装锁定依赖并启动。模型密钥由 DEEPSEEK_API_KEY 环境变量提供；不配置时仍可手工建立任务和使用规则排程。压缩包不包含 node_modules、target、日志或真实密钥。',
 '测试复现时指定另外的测试数据库。后端 Maven verify 生成 Surefire 与 JaCoCo；前端 npm test 和 npm run build 验证工具函数和编译；专用 8088/5182 服务用于 HTTP 黑盒与浏览器场景，固定模型提供方只在测试时启动。',
 '讲解建议：前半分钟介绍问题和成员分工；随后约四十五秒介绍架构与确认边界；约两分半展示创建任务、候选确认、日历和撤销；再用四十五秒展示导入、反馈与测试证据；最后半分钟说明限制和后续计划。录制使用实际界面，不能把模拟测试提供方说成真实模型。'])
table('提交检查',['项目','检查方式'],[['报告','封面姓名学号、目录页码、表前注、实际页数'],['源码','固定提交、独立配置、可重复构建'],['测试','真实执行 JSON/XML 与截图对应'],['视频','约五分钟，解释操作和失败边界'],['压缩包','最终报告哈希与归档内一致']])

# Keep template section geometry; add page fields and update TOC in Word export.
footer=doc.sections[0].footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
settings=doc.settings.element;update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
doc.core_properties.title='学习日程 Agent 系统设计报告';doc.core_properties.author='熊翔、赵欢、黄健华、樊华钰'
output=OUT/'学习日程Agent系统设计报告_最终交付版.docx';doc.save(output)
manifest={'template':str(template),'templateSha256':hashlib.sha256(template.read_bytes()).hexdigest(),'codeCommit':commit,'units':units,'tables':counter,'figures':fig,'bodyCharacters':sum(len(p.text) for p in doc.paragraphs)+sum(len(c.text) for t in doc.tables for r in t.rows for c in r.cells),'memberExcerpts':len(used),'testCount':test_count}
(EV/'report-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf8')
(ROOT/'.runtime/artifact.md').write_text('Teacher template retained; source SHA256 '+manifest['templateSha256']+'\nGeometry and section properties retained. Cover labels filled; body content regions expanded in original chapter order. Original materials unchanged. Table captions and notes precede each added table. Word render and complete page inspection required.\n',encoding='utf8')
print(json.dumps({'file':str(output),'units':len(units),'tables':counter,'figures':fig,'characters':manifest['bodyCharacters']},ensure_ascii=False))
